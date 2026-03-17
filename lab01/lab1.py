import numpy as np
import matplotlib.pyplot as plt
import sounddevice as sd
import soundfile as sf
import scipy.fftpack
from docx import Document
from docx.shared import Inches
from io import BytesIO



data, fs = sf.read('SOUND_INTRO/sound1.wav', dtype='float32')

print(data.dtype)
print(data.shape)

data_L = data[:, 0]
data_R = data[:, 1]

data_mix = (data_R + data_L) / 2

n_samples = data.shape[0]

t = np.arange(n_samples) / fs

sf.write('sound_L.wav', data_L, fs)
sf.write('sound_R.wav', data_R, fs)
sf.write('sound_mix.wav', data_mix, fs)


plt.figure(figsize=(10, 6))
plt.subplot(3, 1, 1)
plt.plot(t,data_L)
plt.title('Left Channel')
plt.subplot(3, 1, 2)
plt.plot(t,data_R)
plt.title('Right Channel')
plt.subplot(3, 1, 3)
plt.plot(t,data_mix)
plt.title('Mixed Channel')
plt.tight_layout()
# plt.show()


data, fs = sf.read('SIN/sin_440Hz.wav', dtype=np.int32)

plt.figure()
plt.subplot(2,1,1)
plt.plot(np.arange(0,data.shape[0])/fs,data)

plt.subplot(2,1,2)
yf = scipy.fftpack.fft(data)
plt.plot(np.arange(0,fs,1.0*fs/(yf.size)),np.abs(yf))
# plt.show()


fsize=2**8

plt.figure()
plt.subplot(2,1,1)
plt.plot(np.arange(0,data.shape[0])/fs,data)
plt.subplot(2,1,2)
yf = scipy.fftpack.fft(data,fsize)
plt.plot(np.arange(0,fs/2,fs/fsize),20*np.log10( np.abs(yf[:fsize//2])))
# plt.show()


#ZAD2 dzwiek

def plotAudio(Signal,Fs,axs,TimeMargin=[0,0.02],fsize = 2 ** 8):

    n_samples = len(Signal)
    t = np.arange(n_samples) / Fs

    axs[0].plot(t, Signal)
    axs[0].set_title("Sygnał w czasie")
    axs[0].set_xlabel("Czas [s]")
    axs[0].set_ylabel("Amplituda")
    axs[0].set_xlim(TimeMargin)
    axs[0].grid(True)

    yf = scipy.fftpack.fft(Signal, fsize)
    amplitude_spec = np.abs(yf[:fsize // 2])
    xf = np.linspace(0.0, Fs / 2, fsize // 2)


    axs[1].plot(np.arange(0, Fs / 2, Fs / fsize), 20 * np.log10(amplitude_spec+ np.finfo(float).eps))
    axs[1].set_title("Widmo amplitudowe 1/2")
    axs[1].set_xlabel("Częstotliwość [Hz]")
    axs[1].set_ylabel("Poziom [dB]")
    axs[1].grid(True)


    idx_max = np.argmax(amplitude_spec)
    freq_max = xf[idx_max]
    amp_max = amplitude_spec[idx_max]

    return freq_max, amp_max




document = Document()
document.add_heading('Wartości dźwięku sinusoidalnego w zależności od fsize', 0)  # tworzenie nagłówków druga wartość to poziom nagłówka

files = ['SIN/sin_60Hz.wav', 'SIN/sin_440Hz.wav', 'SIN/sin_8000Hz.wav']
Margins = [[0, 0.02], [0.133, 0.155]]
fsize=[2**8,2**12,2**16]
for file in files:
    data, fs = sf.read(file, dtype='float32')
    document.add_heading('Plik - {}'.format(file), 2)
    for f in fsize:
        document.add_heading('Rozmiar fsize: {}'.format(f), 3)
        fig, axs = plt.subplots(2, 1, figsize=(10, 7))  # tworzenie plota
        f_max, a_max = plotAudio(data, fs,axs,fsize=f)

    ############################################################
    # Tu wykonujesz jakieś funkcje i rysujesz wykresy
    ############################################################

        fig.suptitle("Plik: {} | fsize: {}".format(file, f)) # Tytuł wykresu
        fig.tight_layout(pad=1.5)  # poprawa czytelności
        memfile = BytesIO()  # tworzenie bufora
        fig.savefig(memfile)  # z zapis do bufora

        document.add_picture(memfile, width=Inches(6))  # dodanie obrazu z bufora do pliku

        memfile.close()
        plt.close(fig)
        ############################################################
        # Tu dodajesz dane tekstowe - wartosci, wyjscie funkcji ect.
        document.add_paragraph(f'Częstotliwość maksymalna: {f_max:.2f} Hz')
        document.add_paragraph(f'Amplituda maksymalna: {a_max:.4f}')
        ############################################################

document.save('report.docx')  # zapis do pliku

