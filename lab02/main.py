import numpy as np
import matplotlib.pyplot as plt
import cv2
import pandas as pd
from matplotlib.pyplot import subplot
from docx import Document
from docx.shared import Inches
from io import BytesIO

img1 = plt.imread('IMG_INTRO/A1.png')
print(img1.dtype)
print(img1.shape)
print(np.min(img1),np.max(img1))

img2 = plt.imread('IMG_INTRO/A2.jpg')
print(img2.dtype)
print(img2.shape)
print(np.min(img2),np.max(img2))

def imgToUInt8(img):
    if np.issubdtype(img.dtype, np.floating):
        img = (img * 255).astype('uint8')
    return img


def imgToFloat(img):
    if np.issubdtype(img.dtype, np.unsignedinteger):
        img =  img.astype(np.float32) / 255.0
    return img

#
# img = plt.imread('IMG_INTRO/A4.jpg')
# plt.imshow(img)
# plt.show()
#
# R=img[:,:,0]
# plt.imshow(R)
# plt.show()
#
# plt.imshow(R, cmap=plt.cm.gray, vmin=0, vmax=255)
# plt.show()
#
# G=img[:,:,1]
# B=img[:,:,2]
#
# Y2=0.2126*R + 0.7152*G + 0.0722*B
# plt.imshow(Y2)
# plt.show()
#
# img_BGR = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
# plt.imshow(img_BGR)
# plt.show()
#

img_zad2 = plt.imread('IMG_INTRO/B01.png')


def plot(img):
    orginal = imgToFloat(img)

    fig = plt.figure(figsize=(20, 20))

    subplot(3, 3, 1)
    plt.imshow(orginal)

    subplot(3, 3, 2)
    kopia1 = orginal.copy()
    Y1 = 0.299 * kopia1[:, :, 0] + 0.587 * kopia1[:, :, 1] + 0.114 * kopia1[:, :, 2]
    plt.imshow(Y1, cmap=plt.cm.gray, vmin=0, vmax=1)

    subplot(3, 3, 3)
    kopia2 = orginal.copy()
    Y2 = 0.2126 * kopia2[:, :, 0] + 0.7152 * kopia2[:, :, 1] + 0.0722 * kopia2[:, :, 2]
    plt.imshow(Y2, cmap=plt.cm.gray, vmin=0, vmax=1)

    plt.subplot(3, 3, 4)
    plt.imshow(orginal[:, :, 0],cmap='gray', vmin=0, vmax=1)
    plt.title("Kanał R")

    plt.subplot(3, 3, 5)
    plt.imshow(orginal[:, :, 1],cmap='gray', vmin=0, vmax=1)
    plt.title("Kanał G")

    plt.subplot(3, 3, 6)
    plt.imshow(orginal[:, :, 2],cmap='gray', vmin=0, vmax=1)
    plt.title("Kanał B")

    plt.subplot(3, 3, 7)
    R_only = orginal.copy()
    R_only[:, :, 1] = 0
    R_only[:, :, 2] = 0
    plt.imshow(R_only)

    plt.subplot(3, 3, 8)
    G_only = orginal.copy()
    G_only[:, :, 0] = 0
    G_only[:, :, 2] = 0
    plt.imshow(G_only)

    plt.subplot(3, 3, 9)
    B_only = orginal.copy()
    B_only[:, :, 0] = 0
    B_only[:, :, 1] = 0
    plt.imshow(B_only)

    return fig

# plot(img_zad2,'output')


fragments = [
    [100, 100, 300, 300],
    [440, 860, 640, 1060],
    [300, 1400, 500, 1600],
    [800, 1600, 1000, 1800],
    [500, 500, 700, 700],
    [100, 1600, 300, 1800],
    [700, 100, 900, 300]
]

df = pd.DataFrame(data={
    'Filename': ['IMG_INTRO/B01.png'],
    'Grayscale': [False],
    'Fragments': [fragments]
})

doc = Document()
doc.add_heading('Raport z analizy fragmentów obrazów', 0)

for index, row in df.iterrows():
    img = plt.imread(row['Filename'])
    doc.add_heading(f'Plik: {row["Filename"]}', level=1)

    if row['Fragments'] is not None:
        for i, f in enumerate(row['Fragments']):
            fragment = img[f[0]:f[2], f[1]:f[3]].copy()
            fig = plot(fragment)

            tytul = f"Plik: {row['Filename']} | Fragment: {i} {f}"
            fig.suptitle(tytul, fontsize=16)
            fig.tight_layout(pad=3.0)

            memfile = BytesIO()
            fig.savefig(memfile, format='png')
            memfile.seek(0)

            doc.add_heading(f'Fragment {i + 1} (wspolrzedne: {f})', level=2)

            doc.add_picture(memfile, width=Inches(6))

            doc.add_paragraph(f"Analiza RGB dla wycinka {f}.")

            memfile.close()
            plt.close(fig)

doc.save('raport2.docx')