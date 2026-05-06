import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
import cv2
import os
import scipy.fftpack
from docx import Document
from docx.shared import Inches
from io import BytesIO
import lab06.main as rle

##########################################
### Settings #############################
##########################################

Test = False

OutputRaportFile = "raport08.docx"

Chroma_options = ["4:4:4", "4:2:2"]
Quant_options = [True, False]

QY = np.array([
    [16, 11, 10, 16, 24, 40, 51, 61],
    [12, 12, 14, 19, 26, 58, 60, 55],
    [14, 13, 16, 24, 40, 57, 69, 56],
    [14, 17, 22, 29, 51, 87, 80, 62],
    [18, 22, 37, 56, 68, 109, 103, 77],
    [24, 36, 55, 64, 81, 104, 113, 92],
    [49, 64, 78, 87, 103, 121, 120, 101],
    [72, 92, 95, 98, 112, 100, 103, 99],
])

QC = np.array([
    [17, 18, 24, 47, 99, 99, 99, 99],
    [18, 21, 26, 66, 99, 99, 99, 99],
    [24, 26, 56, 99, 99, 99, 99, 99],
    [47, 66, 99, 99, 99, 99, 99, 99],
    [99, 99, 99, 99, 99, 99, 99, 99],
    [99, 99, 99, 99, 99, 99, 99, 99],
    [99, 99, 99, 99, 99, 99, 99, 99],
    [99, 99, 99, 99, 99, 99, 99, 99],
])

QN = np.ones((8, 8))

##########################################
### Data Set #############################
##########################################

ImgDir = r'.'  # Address of folder with files (do nor delete `r``)

Images = [
    {
        "Filename": "bruno-guerrero-Ugb9Gaccxys-unsplash.jpg",
        "ROIs": [
            [400, 400, 128, 128],
            [1200, 2400, 128, 128],
            [2000, 3000, 128, 128],
            [3200, 4800, 128, 128],
        ]
    },
    {
        "Filename": "fedor-PtW4RywQV4s-unsplash.jpg",
        "ROIs": [
            [800, 400, 128, 128],
            [2400, 2800, 128, 128],
            [480, 2160, 128, 128],
            [1200, 1920, 128, 128],
        ]
    },
    {
        "Filename": "gian-gomez-rYB1r1MoOXc-unsplash.jpg",
        "ROIs": [
            [400, 400, 128, 128],
            [2400, 800, 128, 128],
            [640, 3200, 128, 128],
            [1600, 2800, 128, 128],
        ]
    },
    {
        "Filename": "oskar-smethurst-B1GtwanCbiw-unsplash.jpg",
        "ROIs": [
            [320, 800, 128, 128],
            [1200, 2400, 128, 128],
            [800, 1840, 128, 128],
            [960, 1264, 128, 128],
        ]
    }
]


##########################################
### Functions to  ########################
##########################################

class JPEG_class:
    Y = np.array([])
    Cb = np.array([])
    Cr = np.array([])
    ChromaRatio = "4:4:4"
    QY = np.ones((8, 8))
    QC = np.ones((8, 8))
    shape = (0, 0, 3)
    compressed_Y = np.array([])
    compressed_Cr = np.array([])
    compressed_Cb = np.array([])


def dct2(a):
    return scipy.fftpack.dct(scipy.fftpack.dct(a.astype(float), axis=0, norm='ortho'), axis=1, norm='ortho')


def idct2(a):
    return scipy.fftpack.idct(scipy.fftpack.idct(a.astype(float), axis=0, norm='ortho'), axis=1, norm='ortho')


def zigzag(A):
    template = np.array([
        [0, 1, 5, 6, 14, 15, 27, 28],
        [2, 4, 7, 13, 16, 26, 29, 42],
        [3, 8, 12, 17, 25, 30, 41, 43],
        [9, 11, 18, 24, 31, 40, 44, 53],
        [10, 19, 23, 32, 39, 45, 52, 54],
        [20, 22, 33, 38, 46, 51, 55, 60],
        [21, 34, 37, 47, 50, 56, 59, 61],
        [35, 36, 48, 49, 57, 58, 62, 63],
    ])
    if len(A.shape) == 1:
        B = np.zeros((8, 8))
        for r in range(0, 8):
            for c in range(0, 8):
                B[r, c] = A[template[r, c]]
    else:
        B = np.zeros((64,))
        for r in range(0, 8):
            for c in range(0, 8):
                B[template[r, c]] = A[r, c]
    return B


def CompressBlock(block, Q):
    block_centered = block.astype(float) - 128.0
    dct_block = dct2(block_centered)

    quantized_block = np.round(dct_block / Q).astype(int)

    vector = zigzag(quantized_block)
    return vector


def DecompressBlock(vector, Q):
    block = zigzag(vector)

    dequantized_block = block * Q

    idct_block = idct2(dequantized_block)
    block_restored = idct_block + 128.0
    return block_restored


def CompressLayer(L, Q):
    S = np.array([])
    for w in range(0, L.shape[0], 8):
        for k in range(0, L.shape[1], 8):
            block = L[w:(w + 8), k:(k + 8)]
            S = np.append(S, CompressBlock(block, Q))
    return S


def DecompressLayer(S, shape, Q):
    L = np.zeros(shape)
    for idx, i in enumerate(range(0, S.shape[0], 64)):
        vector = S[i:(i + 64)]
        m = L.shape[1] / 8
        k = int((idx % m) * 8)
        w = int((idx // m) * 8)
        L[w:(w + 8), k:(k + 8)] = DecompressBlock(vector, Q)
    return L


def CompressJPEG(RGB, Ratio="4:4:4", QY=QN, QC=QN):
    JPEG = JPEG_class()
    JPEG.shape = RGB.shape
    JPEG.ChromaRatio = Ratio
    JPEG.QY = QY
    JPEG.QC = QC

    YCrCb = cv2.cvtColor(RGB, cv2.COLOR_RGB2YCrCb).astype(int)

    JPEG.Y = YCrCb[:, :, 0]
    JPEG.Cr = YCrCb[:, :, 1]
    JPEG.Cb = YCrCb[:, :, 2]

    if Ratio == "4:2:2":
        JPEG.Cr = JPEG.Cr[:, ::2]
        JPEG.Cb = JPEG.Cb[:, ::2]
    elif Ratio == "4:2:0":
        pass
    else:  # default "4:4:4"
        pass

    JPEG.Y = CompressLayer(JPEG.Y, JPEG.QY)
    JPEG.Cr = CompressLayer(JPEG.Cr, JPEG.QC)
    JPEG.Cb = CompressLayer(JPEG.Cb, JPEG.QC)

    JPEG.compressed_Y = rle.rle_encoder(JPEG.Y)
    JPEG.compressed_Cr = rle.rle_encoder(JPEG.Cr)
    JPEG.compressed_Cb = rle.rle_encoder(JPEG.Cb)

    print(f"\n[Chroma: {Ratio}] Kompresja RLE:")
    print(f"Y:  {len(JPEG.Y)} -> {len(JPEG.compressed_Y)} elementów")
    print(f"Cr: {len(JPEG.Cr)} -> {len(JPEG.compressed_Cr)} elementów")
    print(f"Cb: {len(JPEG.Cb)} -> {len(JPEG.compressed_Cb)} elementów")

    return JPEG


def DecompressJPEG(JPEG):
    layer_shape_Y = (JPEG.shape[0], JPEG.shape[1])

    if JPEG.ChromaRatio == "4:2:2":
        layer_shape_C = (JPEG.shape[0], int(JPEG.shape[1] / 2))
    else:
        layer_shape_C = (JPEG.shape[0], JPEG.shape[1])

    Y_comp = rle.rle_decoder(JPEG.compressed_Y)
    Cr_comp = rle.rle_decoder(JPEG.compressed_Cr)
    Cb_comp = rle.rle_decoder(JPEG.compressed_Cb)

    Y = DecompressLayer(Y_comp, layer_shape_Y, JPEG.QY)
    Cr = DecompressLayer(Cr_comp, layer_shape_C, JPEG.QC)
    Cb = DecompressLayer(Cb_comp, layer_shape_C, JPEG.QC)

    if JPEG.ChromaRatio == "4:2:2":
        Cr = np.repeat(Cr, 2, axis=1)
        Cb = np.repeat(Cb, 2, axis=1)
    elif JPEG.ChromaRatio == "4:2:0":
        pass
    else:  # default "4:4:4"
        pass

    YCrCb = np.zeros(JPEG.shape, dtype=int)
    YCrCb[:, :, 0] = Y
    YCrCb[:, :, 1] = Cr
    YCrCb[:, :, 2] = Cb

    RGB = cv2.cvtColor(YCrCb.astype(np.uint8), cv2.COLOR_YCrCb2RGB)

    return RGB


##########################################
### Main Program  ########################
##########################################

def plot_comparisone(counter, OG, Decomp, figsize=(5, 8)):
    fig, axs = plt.subplots(4, 2, num=counter, sharex=True, sharey=True, figsize=figsize)
    axs[0, 0].imshow(OG)  # RGB
    PRZED_YCrCb = cv2.cvtColor(OG, cv2.COLOR_RGB2YCrCb)
    axs[1, 0].imshow(PRZED_YCrCb[:, :, 0], cmap='gray')
    axs[2, 0].imshow(PRZED_YCrCb[:, :, 1], cmap='gray')
    axs[3, 0].imshow(PRZED_YCrCb[:, :, 2], cmap='gray')

    axs[0, 0].set_title("Oryginał")
    axs[1, 0].set_title("Y")
    axs[2, 0].set_title("Cr")
    axs[3, 0].set_title("Cb")

    axs[0, 1].imshow(Decomp)  # RGB
    PO_YCrCb = cv2.cvtColor(Decomp, cv2.COLOR_RGB2YCrCb)
    axs[1, 1].imshow(PO_YCrCb[:, :, 0], cmap='gray')
    axs[2, 1].imshow(PO_YCrCb[:, :, 1], cmap='gray')
    axs[3, 1].imshow(PO_YCrCb[:, :, 2], cmap='gray')

    axs[0, 1].set_title("Po kompresji")
    axs[1, 1].set_title("Y")
    axs[2, 1].set_title("Cr")
    axs[3, 1].set_title("Cb")

    for ax in axs.flatten():
        ax.set_axis_off()

    return fig


if Test:
    img = plt.imread(os.path.join(ImgDir, Images[0]["Filename"]))
    ROI = Images[0]["ROIs"][0]
    fragment = img[ROI[1]:ROI[1] + ROI[3], ROI[0]:ROI[0] + ROI[2]]
    Counter = 1
    for Chroma in Chroma_options:
        for Quant in Quant_options:
            if Quant:
                tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QY, QC=QC)
            else:
                tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QN, QC=QN)

            New_Fragment = DecompressJPEG(tJPEG)

            f = plot_comparisone(Counter, fragment, New_Fragment)
            f.suptitle("Plik: {} Chroma: {} Kwantyzacja: {}".format(Images[0]["Filename"], Chroma, Quant))
            Counter += 1
    plt.show()

else:
    # generate raport
    document = Document()
    document.add_heading('Report', 0)
    document.add_paragraph("Autor: Norbert Świstak")
    document.add_section()
    document.add_heading("Fragmenty wygenerowane na podstawie działania funkcji", 1)
    Counter = 1
    for file_dict in Images:
        filename = file_dict['Filename']
        img = plt.imread(os.path.join(ImgDir, filename))
        for ROI in file_dict['ROIs']:
            fragment = img[ROI[1]:ROI[1] + ROI[3], ROI[0]:ROI[0] + ROI[2]]

            for Chroma in Chroma_options:
                for Quant in Quant_options:
                    if Quant:
                        tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QY, QC=QC)
                    else:
                        tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QN, QC=QN)

                    New_Fragment = DecompressJPEG(tJPEG)

                    f = plot_comparisone(Counter, fragment, New_Fragment)
                    f.suptitle("Plik: {} Chroma: {} Kwantyzacja: {}".format(filename, Chroma, Quant))
                    memfile = BytesIO()
                    f.savefig(memfile)
                    document.add_picture(memfile, width=Inches(6))  # set document size
                    memfile.close()
                    f.clf()

                    document.add_paragraph(f"Parametry: Chroma = {Chroma}, Kwantyzacja = {Quant}")
                    document.add_paragraph(
                        f"Warstwa Y:  {len(tJPEG.Y)} -> {len(tJPEG.compressed_Y)} elementów "
                        f"({(len(tJPEG.compressed_Y) / max(1, len(tJPEG.Y))) * 100:.2f}%)"
                    )
                    document.add_paragraph(
                        f"Warstwa Cr: {len(tJPEG.Cr)} -> {len(tJPEG.compressed_Cr)} elementów "
                        f"({(len(tJPEG.compressed_Cr) / max(1, len(tJPEG.Cr))) * 100:.2f}%)"
                    )
                    document.add_paragraph(
                        f"Warstwa Cb: {len(tJPEG.Cb)} -> {len(tJPEG.compressed_Cb)} elementów "
                        f"({(len(tJPEG.compressed_Cb) / max(1, len(tJPEG.Cb))) * 100:.2f}%)"
                    )
                    document.add_paragraph("-" * 40)
            Counter += 1
    document.add_section()
    document.add_heading("Podsumowanie i wnioski", 1)
    document.add_paragraph("Tu proszę zebrać wszystkie obserwacje na podstawie powyższych wykresów i napisać wnioski.")
    document.save(OutputRaportFile)